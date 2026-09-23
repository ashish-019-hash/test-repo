#!/usr/bin/env python
"""Render `tests/fixtures/telecom_spec.md` to PDF and DOCX, plus a blank OCR-test PDF.

Re-runnable: `python scripts/make_fixtures.py` regenerates the committed fixtures in
place. `generate(md_path, out_dir)` is used directly by
`tests/unit/test_fixtures_regeneration.py` to assert that regeneration is stable.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, TableStyle
from reportlab.platypus import Table as RLTable

FIXTURES_DIR = Path(__file__).resolve().parent.parent / "tests" / "fixtures"
AUTHOR = "Doc Extractor Fixtures"

HEADING_SIZES = {1: 20, 2: 16, 3: 13}
BODY_SIZE = 10
FOOTER_SIZE = 8
_PAGE_WIDTH, _PAGE_HEIGHT = LETTER
_MARGIN = 0.75 * inch
_AVAIL_WIDTH = _PAGE_WIDTH - 2 * _MARGIN


@dataclass
class MdElement:
    kind: str  # "heading" | "paragraph" | "table"
    level: int | None = None
    text: str = ""
    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


def parse_markdown(content: str) -> list[MdElement]:
    """Minimal parser for the fixture markdown: headings, pipe tables, blank-line paragraphs."""
    lines = content.splitlines()
    elements: list[MdElement] = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        stripped = line.lstrip()
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            elements.append(MdElement(kind="heading", level=level, text=stripped[level:].strip()))
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [[c.strip() for c in tl.strip("|").split("|")] for tl in table_lines]
            rows = [r for r in rows if not all(re.fullmatch(r"-{2,}", c) for c in r)]
            header, *data_rows = rows
            elements.append(MdElement(kind="table", header=header, rows=data_rows))
            continue
        para_lines = [line.strip()]
        i += 1
        while i < n and lines[i].strip() and not lines[i].lstrip().startswith(("#", "|")):
            para_lines.append(lines[i].strip())
            i += 1
        elements.append(MdElement(kind="paragraph", text=" ".join(para_lines)))
    return elements


# ---------------------------------------------------------------------------------------- PDF
class _NumberedCanvas(Canvas):
    """Draws 'Page N of M' on every page; M is only known once the whole doc is built."""

    def __init__(self, *args: object, **kwargs: object) -> None:
        super().__init__(*args, **kwargs)
        self._saved_page_states: list[dict[str, object]] = []

    def showPage(self) -> None:  # noqa: N802 (reportlab API name)
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        total = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self._draw_footer(total)
            super().showPage()
        super().save()

    def _draw_footer(self, total: int) -> None:
        self.setFont("Helvetica", FOOTER_SIZE)
        self.drawCentredString(_PAGE_WIDTH / 2.0, 0.4 * inch, f"Page {self._pageNumber} of {total}")


def _heading_style(level: int) -> ParagraphStyle:
    size = HEADING_SIZES.get(level, HEADING_SIZES[3])
    return ParagraphStyle(
        name=f"heading{level}",
        fontName="Helvetica-Bold",
        fontSize=size,
        leading=size * 1.2,
        spaceBefore=12,
        spaceAfter=6,
    )


_BODY_STYLE = ParagraphStyle(
    name="body", fontName="Helvetica", fontSize=BODY_SIZE, leading=BODY_SIZE * 1.3, spaceAfter=6
)
_CELL_STYLE = ParagraphStyle(name="cell", fontName="Helvetica", fontSize=8, leading=9.5)
_CELL_HEADER_STYLE = ParagraphStyle(name="cell-header", fontName="Helvetica-Bold", fontSize=8, leading=9.5)


def _col_widths(header: list[str], rows: list[list[str]]) -> list[float]:
    """Each column must be at least as wide as its longest single word (no font size 8) so
    reportlab never force-splits a word mid-character; remaining width is distributed by
    total column content length so long free-text columns (e.g. Description) get more room.
    """
    all_rows = [header, *rows]
    ncols = len(header)
    min_widths = []
    for c in range(ncols):
        words = [w for r in all_rows for w in (r[c].split() or [""])]
        longest = max((stringWidth(w, "Helvetica", 8) for w in words), default=0.0)
        min_widths.append(longest + 8.0)  # cell padding
    weights = [sum(len(r[c]) for r in all_rows) or 1 for c in range(ncols)]
    total_min = sum(min_widths)
    extra = max(_AVAIL_WIDTH - total_min, 0.0)
    total_weight = sum(weights)
    return [
        min_widths[c] + (extra * weights[c] / total_weight if total_weight else 0.0) for c in range(ncols)
    ]


def _table_flowable(header: list[str], rows: list[list[str]]) -> RLTable:
    col_widths = _col_widths(header, rows)
    data = [[Paragraph(c, _CELL_HEADER_STYLE) for c in header]]
    for row in rows:
        data.append([Paragraph(c if c else "&nbsp;", _CELL_STYLE) for c in row])
    table = RLTable(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.6, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def render_pdf(elements: list[MdElement], out_path: Path, title: str, author: str) -> Path:
    flowables: list[object] = []
    for el in elements:
        if el.kind == "heading":
            flowables.append(Paragraph(el.text, _heading_style(el.level or 3)))
        elif el.kind == "paragraph":
            flowables.append(Paragraph(el.text, _BODY_STYLE))
        elif el.kind == "table":
            flowables.append(_table_flowable(el.header, el.rows))
            flowables.append(Spacer(1, 10))
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=LETTER,
        title=title,
        author=author,
        invariant=1,
        topMargin=_MARGIN,
        bottomMargin=_MARGIN,
        leftMargin=_MARGIN,
        rightMargin=_MARGIN,
    )
    doc.build(flowables, canvasmaker=_NumberedCanvas)
    return out_path


def render_blank_pdf(out_path: Path) -> Path:
    """One page, no extractable text at all (used to exercise OCR-detection)."""
    canvas = Canvas(str(out_path), pagesize=LETTER, invariant=1)
    canvas.showPage()
    canvas.save()
    return out_path


# --------------------------------------------------------------------------------------- DOCX
def render_docx(elements: list[MdElement], out_path: Path, title: str, author: str) -> Path:
    doc = DocxDocument()
    doc.core_properties.title = title
    doc.core_properties.author = author
    for el in elements:
        if el.kind == "heading":
            level = el.level or 3
            style = "Title" if level == 1 else f"Heading {level - 1}"
            doc.add_paragraph(el.text, style=style)
        elif el.kind == "paragraph":
            doc.add_paragraph(el.text)
        elif el.kind == "table":
            ncols = len(el.header)
            table = doc.add_table(rows=1, cols=ncols)
            table.style = "Table Grid"
            for i, cell_text in enumerate(el.header):
                table.rows[0].cells[i].text = cell_text
            for row in el.rows:
                cells = table.add_row().cells
                for i, cell_text in enumerate(row):
                    cells[i].text = cell_text
    doc.save(str(out_path))
    _normalize_zip_timestamps(out_path)
    return out_path


_FIXED_ZIP_DATE = (2013, 12, 23, 23, 15, 0)


def _normalize_zip_timestamps(path: Path) -> None:
    """DOCX is a zip; python-docx stamps each entry with the current wall-clock time,
    which would make the fixture non-reproducible across runs. Pin every entry to the
    same fixed date (matching the docProps/core.xml created/modified dates already baked
    into python-docx's default template) so the output is byte-stable.
    """
    data = path.read_bytes()
    buf = io.BytesIO(data)
    with zipfile.ZipFile(buf) as src:
        entries = [(info, src.read(info.filename)) for info in src.infolist()]

    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, blob in entries:
            info.date_time = _FIXED_ZIP_DATE
            dst.writestr(info, blob)
    path.write_bytes(out_buf.getvalue())


# ---------------------------------------------------------------------------------- entrypoint
def generate(md_path: Path, out_dir: Path) -> dict[str, Path]:
    content = md_path.read_text(encoding="utf-8")
    elements = parse_markdown(content)
    title = next((e.text for e in elements if e.kind == "heading" and e.level == 1), "Telecom Spec")
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = render_pdf(elements, out_dir / "telecom_spec.pdf", title=title, author=AUTHOR)
    docx_path = render_docx(elements, out_dir / "telecom_spec.docx", title=title, author=AUTHOR)
    blank_path = render_blank_pdf(out_dir / "blank_page.pdf")
    return {"pdf": pdf_path, "docx": docx_path, "blank_pdf": blank_path}


def main() -> None:
    paths = generate(FIXTURES_DIR / "telecom_spec.md", FIXTURES_DIR)
    for kind, path in paths.items():
        print(f"{kind}: {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
